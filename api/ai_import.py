"""Shared AI-powered import: turn ANY uploaded file (Excel, CSV, PDF, Word,
image/photo, plain text) into either clean text or structured rows, so every
"Import" button in the app behaves the same — accept anything, let the model
read it, and return a preview the coach confirms before committing.

The heavy lifting is one call to Claude with the file as native content blocks
(images/PDFs as vision/document blocks; spreadsheets/docs as extracted text).
"""
from __future__ import annotations

import base64
import io
import json
import re
from typing import Any
from .ai_models import SONNET, text_of


def file_to_content_blocks(data: bytes, filename: str, content_type: str | None) -> tuple[list[dict], bool]:
    """Turn raw upload bytes into Claude content blocks. Returns (blocks, is_text)
    where is_text is True when the content is plain extracted text (no vision/PDF
    block needed). Best-effort — never raises."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    # Images → vision block
    if ctype.startswith("image/") or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".heic", ".heif")):
        media = ("image/png" if name.endswith(".png") else
                 "image/webp" if name.endswith(".webp") else
                 "image/gif" if name.endswith(".gif") else "image/jpeg")
        return [{"type": "image", "source": {"type": "base64", "media_type": media, "data": base64.b64encode(data).decode()}}], False

    # PDF → native document block
    if ctype == "application/pdf" or name.endswith(".pdf"):
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()}}], False

    # Excel → text grid
    if name.endswith((".xlsx", ".xls")) or "spreadsheet" in ctype or "excel" in ctype:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return [{"type": "text", "text": "\n".join(lines)[:20000]}], True
        except Exception:
            pass

    # Word
    if name.endswith(".docx") or "word" in ctype:
        try:
            import docx
            d = docx.Document(io.BytesIO(data))
            return [{"type": "text", "text": "\n".join(p.text for p in d.paragraphs if p.text.strip())[:20000]}], True
        except Exception:
            pass

    # Plain text / CSV / anything decodable
    try:
        return [{"type": "text", "text": data.decode("utf-8", errors="ignore")[:20000]}], True
    except Exception:
        return [], True


def _client():
    import anthropic
    return anthropic.Anthropic()


def _extract_json(text: str) -> Any:
    """Pull the first JSON object/array out of a model response."""
    text = text.strip()
    # Strip code fences
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.MULTILINE).strip()
    try:
        return json.loads(text)
    except Exception:
        pass
    m = re.search(r"(\{.*\}|\[.*\])", text, re.DOTALL)
    if m:
        try:
            return json.loads(m.group(1))
        except Exception:
            return None
    return None


def ai_extract_json(data: bytes, filename: str, content_type: str | None, instruction: str) -> Any:
    """Feed the file to Claude with `instruction` (which must ask for JSON only)
    and return the parsed JSON, or None."""
    import os
    if not os.environ.get("ANTHROPIC_API_KEY"):
        raise RuntimeError("ANTHROPIC_API_KEY is not configured on the server.")
    blocks, _ = file_to_content_blocks(data, filename, content_type)
    if not blocks:
        return None
    content = [{"type": "text", "text": instruction}] + blocks
    resp = _client().messages.create(
        model=SONNET, max_tokens=8000,
        messages=[{"role": "user", "content": content}],
    )
    return _extract_json(text_of(resp))


def ai_extract_text(data: bytes, filename: str, content_type: str | None, purpose: str) -> str:
    """Return clean text from the file. If it's already text-extractable we return
    it directly; images/PDFs are transcribed by the model, focused on `purpose`."""
    blocks, is_text = file_to_content_blocks(data, filename, content_type)
    if not blocks:
        return ""
    if is_text:
        # Already have the text — return it as-is (fast, no AI needed).
        return blocks[0].get("text", "").strip()
    import os
    if not os.environ.get("ANTHROPIC_API_KEY"):
        raise RuntimeError("ANTHROPIC_API_KEY is not configured on the server.")
    instruction = (
        f"Read the attached file and extract its text content for use as {purpose}. "
        "Transcribe the meaningful content faithfully and cleanly (no commentary, no "
        "markdown fences). If it is a handwritten or photographed note, transcribe what "
        "it says. Return ONLY the extracted text."
    )
    content = [{"type": "text", "text": instruction}] + blocks
    resp = _client().messages.create(
        model=SONNET, max_tokens=6000,
        messages=[{"role": "user", "content": content}],
    )
    return text_of(resp).strip()


# ── Context-specific structured extractors ────────────────────────────────────

ROSTER_INSTRUCTION = (
    "You are extracting a basketball ROSTER from the attached file (it may be a "
    "spreadsheet, PDF, Word doc, photo of a roster, or text). Return ONLY a JSON "
    "object of this exact shape:\n"
    '{"players": [{"name": "", "jersey_number": "", "position": "", "height": "", '
    '"wingspan": "", "weight": "", "school_name": "", "competition_level": ""}]}\n'
    "Rules: one entry per player. name is required — skip rows without a name. Use "
    "empty string for any field not present. jersey_number as a string (e.g. \"23\"). "
    "height like 6'2\". Do not invent players or data. Return JSON only, no prose."
)

GAME_STATS_INSTRUCTION = (
    "You are extracting a basketball BOX SCORE / stat sheet from the attached file "
    "(spreadsheet, PDF, photo, or text). Return ONLY a JSON object of this shape:\n"
    '{"players": [{"player_name": "", "team_name": "", "is_opponent": false, '
    '"stats": {"2 FG Made": 0, "2 FG Missed": 0, "3 FG Made": 0, "3 FG Missed": 0, '
    '"FT Made": 0, "FT Missed": 0, "Off. Reb": 0, '
    '"Def. Reb": 0, "Assists": 0, "Steal": 0, "Blocked Shot": 0, "Turnover": 0, '
    '"Foul Against": 0}}]}\n'
    "Rules: one entry per player. Map the file's columns to these stat names (e.g. "
    "PTS/REB/AST/STL/BLK/TO/PF, made field goals, threes, free throws).\n"
    "MISSES MATTER AS MUCH AS MAKES. A box score states attempts as well as makes "
    "— '9-18', or an FGA column beside FGM. A miss is an attempt that was not a "
    "make, so subtract: 18 attempted and 9 made is 9 missed. Do this for field "
    "goals, threes and free throws. Without them no shooting percentage can be "
    "worked out at all, which is the single most common thing a coach asks of a "
    "box score.\n"
    "FG/FGM/FGA COVER ALL FIELD GOALS, THREES INCLUDED. '2 FG Made' must be the "
    "TWO-POINT makes only: subtract the threes. A 9-of-18 line containing three "
    "threes is 6 two-point makes and 3 three-point makes, not 9 and 3.\n"
    "TEAM_NAME IS THE HEADING THE FILE PUTS THE PLAYER UNDER — copy it exactly, "
    "even if you are unsure which side it is. A sheet almost always names both "
    "teams, and that name is what lets a coach say which side it belongs to; "
    "guessing wrong is recoverable, leaving it blank is not.\n"
    "Set is_opponent true for players clearly on the opposing team. Omit any stat "
    "the file does not state rather than guessing it — a zero here reads as a "
    "real zero. Skip team-total rows. Return JSON only, no prose."
)


# One pass per file, for everything a game file might hold.
#
# Asking three times — "is this a box score", "is this a play-by-play", "are
# there shot locations" — costs three reads of the same page and still misses
# the common case, which is one export carrying several of them. This asks once
# and takes whatever is there.
#
# The rule throughout is the same one the rest of the app follows: state what
# the file states, omit what it does not. An invented coordinate or an estimated
# fast-break total reads as measurement, and a coach has no way to tell it from
# the real thing.
GAME_FILE_INSTRUCTION = (
    "You are reading a basketball game file: a box score, a play-by-play, a shot "
    "chart, a team-stats panel, or several of those together. It may be a "
    "spreadsheet, a PDF, a photo of a printed sheet, or a screenshot.\n\n"
    "Return ONLY this JSON, including whichever sections the file actually "
    "contains and omitting the rest:\n"
    '{\n'
    '  "players": [{"player_name": "", "jersey_number": "", "team_name": "", "is_opponent": false, '
    '"MIN": "0:00", "PM": 0, "EFF": 0, '
    '"stats": {"2 FG Made": 0, "2 FG Missed": 0, "3 FG Made": 0, "3 FG Missed": 0, '
    '"FT Made": 0, "FT Missed": 0, "Off. Reb": 0, "Def. Reb": 0, "Assists": 0, '
    '"Steal": 0, "Blocked Shot": 0, "Turnover": 0, "Foul Against": 0}}],\n'
    '  "events": [{"period": 1, "clock": "8:32", "team_name": "", "player_name": "", '
    '"description": "", "points": 0, "home_score": 0, "away_score": 0}],\n'
    '  "shots": [{"team_name": "", "player_name": "", "period": 1, "x": 0.0, "y": 0.0, '
    '"made": true, "points": 2}],\n'
    '  "team_stats": [{"team_name": "", "PTS": 0, "REB": 0, "OREB": 0, "DREB": 0, "AST": 0, '
    '"STL": 0, "BLK": 0, "TO": 0, "PF": 0, "points_off_turnovers": 0, "fast_break_points": 0, '
    '"second_chance_points": 0, "points_in_paint": 0, "bench_points": 0}]\n'
    '}\n\n'
    "PLAYERS — the box score. jersey_number is the squad number printed beside the "
    "name, as text ('00' is not 0); omit it when the sheet does not print one.\n"
    "MIN is minutes played exactly as printed — '27:13' stays '27:13', 27 stays 27. "
    "PM is the plus-minus column and CAN be negative; keep the sign. EFF is the "
    "efficiency or PIR column. These three are not counted stats and belong beside "
    "the name, not in 'stats'. Omit any the sheet does not print rather than "
    "working it out.\n"
    "FG/FGM/FGA cover ALL field goals, threes included, so "
    "'2 FG Made' must have the threes subtracted: a 9-of-18 line containing three "
    "threes is 6 two-point makes and 3 three-point makes. Misses matter as much as "
    "makes — a miss is an attempt that was not a make, so subtract. Without them no "
    "shooting percentage can be worked out at all.\n\n"
    "EVENTS — the play-by-play, in the order the file lists them. 'clock' is the game "
    "clock as printed (MM:SS remaining in the period). 'points' is how many this event "
    "scored, 0 for anything that did not score. home_score/away_score are the running "
    "score AFTER the event, when the file prints them. Include every row, not just the "
    "scoring ones: turnovers, rebounds and fouls are what make the rest readable.\n\n"
    "SHOTS — only when the file gives a LOCATION for each shot. x and y are "
    "percentages of the full court, 0-100, with x running along the length of the "
    "court and y across it. If the file gives coordinates in feet or in its own units, "
    "convert to percentages. If it shows shots on a chart image, read the positions as "
    "accurately as you can from the image. Never invent a location for a shot the file "
    "only counts.\n\n"
    "TEAM_STATS — the TEAM TOTALS row or panel, when the sheet prints one. Give the "
    "totals EXACTLY as printed; do NOT add the player rows up yourself. They are "
    "often higher than the players sum to, because team rebounds and team turnovers "
    "belong to the team rather than to any player, and that difference is the whole "
    "reason this section exists. Points off turnovers, fast-break, second-chance, "
    "paint and bench points cannot be worked out from a box score at all, so give "
    "them only when the file states them.\n"
    "These five are usually NOT in a table. They come as a comparison chart: the "
    "stat's name down the middle — 'Points off Turnovers', 'Fast Break Points', "
    "'Second Chance Points', 'Points in the Paint', 'Points from the Bench' — with "
    "one team's number and bar to the left of it and the other team's to the right. "
    "Read both ends of every row; that is two team_stats rows, not one. A small "
    "triangle or arrow beside a number only marks the higher side and is not part of "
    "the number. The same layout is used for PTS/REB/AST/STL/BLK/TO/PF and for "
    "shooting percentages — read those the same way.\n\n"
    "TEAM_NAME everywhere is the heading the file puts that row under — copy it "
    "exactly even if unsure which side it is. Set is_opponent when the file makes the "
    "sides clear; when it does not, leave it false and let the team name speak.\n\n"
    "WHEN THE FILE DOES NOT NAME THE TEAMS — a chart with two coloured columns, a "
    "sheet headed only HOME and AWAY, two blocks with no headings — do NOT guess a "
    "name and do NOT leave team_name empty. Give each side a short label describing "
    "where it is and what it looks like, exactly as a person would point at it: "
    "\"red, left\" and \"blue, right\", or \"top block\" and \"bottom block\". Use the "
    "SAME label for every row of that side, in every section. A guessed name reads as "
    "a known one and puts a whole team's numbers on the wrong side; a label the coach "
    "recognises lets them say which is which.\n\n"
    "Omit any field the file does not state rather than guessing it: a zero reads as a "
    "real zero, and an invented shot location reads as a measurement. Skip team-total "
    "rows in players. Return JSON only, no prose."
)
